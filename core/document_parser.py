# -*- coding: utf-8 -*-
"""文档与图片内容提取：为“附件生成测试用例”提供统一文本输入。"""

import io
import re
import logging
from pathlib import Path
from typing import Any, Dict

logger = logging.getLogger("DocumentParser")

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}
TEXT_EXTS = {".txt", ".md", ".markdown", ".csv", ".log", ".json", ".xml", ".yaml", ".yml"}
DOCUMENT_EXTS = IMAGE_EXTS | TEXT_EXTS | {".pdf", ".docx", ".xlsx"}
SUPPORTED_EXTS = DOCUMENT_EXTS
MAX_FILE_SIZE = 20 * 1024 * 1024
MAX_TEXT_CHARS = 60000


def _clean_text(text: str) -> str:
    text = (text or "").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_image(raw: bytes) -> str:
    from core.vision_ocr import VisionOCR

    elements, _ = VisionOCR.recognize_elements(raw)
    lines = []
    for item in elements or []:
        text = str(item.get("text") or "").strip()
        if text:
            lines.append(text)
    return "\n".join(lines)


def _extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("PDF 已加密，无法提取文字")
    parts = []
    for page in reader.pages[:60]:
        try:
            parts.append(page.extract_text() or "")
        except Exception as e:
            logger.warning(f"PDF 单页解析失败: {e}")
    text = "\n".join(parts)
    if not text.strip():
        raise ValueError("PDF 中没有可提取的文字（可能是扫描件，请改用截图上传走 OCR）")
    return text


def _extract_docx(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(raw: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    parts = []
    try:
        for sheet in wb.worksheets[:10]:
            parts.append(f"## {sheet.title}")
            for idx, row in enumerate(sheet.iter_rows(values_only=True)):
                if idx >= 2000:
                    parts.append("...（表格行数过多，已截断）")
                    break
                values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if values:
                    parts.append(" | ".join(values))
    finally:
        wb.close()
    return "\n".join(parts)


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_document(filename: str, raw: bytes) -> Dict[str, Any]:
    """提取单个附件的文本内容，返回统一的解析结果。"""
    safe_name = Path(filename or "attachment").name
    ext = Path(safe_name).suffix.lower()

    if not raw:
        raise ValueError("文件内容为空")
    if len(raw) > MAX_FILE_SIZE:
        raise ValueError(f"文件超过 {MAX_FILE_SIZE // (1024 * 1024)}MB 限制")
    if ext not in DOCUMENT_EXTS:
        raise ValueError(f"暂不支持 {ext or '未知'} 类型，请上传 PDF、DOCX、TXT/MD/CSV、XLSX 或图片")

    if ext in IMAGE_EXTS:
        text = _extract_image(raw)
        kind = "image"
        warning = "" if text.strip() else "图片中未识别到文字，请确认图片清晰度"
    elif ext == ".pdf":
        text = _extract_pdf(raw)
        kind = "pdf"
        warning = ""
    elif ext == ".docx":
        text = _extract_docx(raw)
        kind = "docx"
        warning = ""
    elif ext == ".xlsx":
        text = _extract_xlsx(raw)
        kind = "xlsx"
        warning = ""
    else:
        text = _decode_text(raw)
        kind = "text"
        warning = ""

    text = _clean_text(text)
    truncated = len(text) > MAX_TEXT_CHARS
    if truncated:
        text = text[:MAX_TEXT_CHARS] + "\n\n...（内容过长，已截断）"

    return {
        "filename": safe_name,
        "kind": kind,
        "text": text,
        "chars": len(text),
        "truncated": truncated,
        "warning": warning,
    }
